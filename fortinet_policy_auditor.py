#!/usr/bin/env python3
"""Fortinet Policy Auditor

Analiza reglas IPv4 procedentes de CSV, JSON o texto de configuración FortiGate,
extrae inventario desde exportaciones .dat/.tgz de FortiManager, y genera JSON,
CSV y XLSX. El análisis es estático y no modifica dispositivos.
"""
from __future__ import annotations
import argparse, csv, gzip, hashlib, ipaddress, json, re, tarfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Iterable

ANY = "all"

@dataclass
class Rule:
    device: str = "unknown"
    vdom: str = "root"
    policy_id: str = ""
    sequence: int = 0
    name: str = ""
    srcintf: str = ANY
    dstintf: str = ANY
    srcaddr: str = ANY
    dstaddr: str = ANY
    service: str = ANY
    schedule: str = "always"
    action: str = "deny"
    status: str = "enable"
    comments: str = ""

    def normalized(self):
        return {
            "srcintf": norm_tokens(self.srcintf), "dstintf": norm_tokens(self.dstintf),
            "srcaddr": norm_tokens(self.srcaddr), "dstaddr": norm_tokens(self.dstaddr),
            "service": norm_tokens(self.service), "schedule": norm_tokens(self.schedule),
            "action": self.action.strip().lower(), "status": self.status.strip().lower()
        }

FIELDS = [f.name for f in Rule.__dataclass_fields__.values()]


def norm_tokens(v: Any) -> tuple[str, ...]:
    if isinstance(v, list): xs = [str(x) for x in v]
    else: xs = re.split(r"[;,|]", str(v or ANY))
    out = sorted({x.strip().strip('"').lower() for x in xs if x.strip()})
    return tuple(out or [ANY])


def read_json(path: Path) -> list[Rule]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(data, dict): data = data.get("rules", data.get("policies", [data]))
    return [coerce_rule(x, i+1) for i, x in enumerate(data)]


def read_csv(path: Path) -> list[Rule]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return [coerce_rule(r, i+1) for i, r in enumerate(csv.DictReader(f))]


def coerce_rule(d: dict, seq: int) -> Rule:
    aliases = {"id":"policy_id", "policyid":"policy_id", "src":"srcaddr", "dst":"dstaddr",
               "source":"srcaddr", "destination":"dstaddr", "services":"service", "policy_name":"name"}
    clean = {}
    for k,v in d.items():
        key = aliases.get(str(k).strip().lower().replace("-","_").replace(" ","_"), str(k).strip().lower())
        if key in FIELDS: clean[key] = v
    clean["sequence"] = int(clean.get("sequence") or seq)
    clean["policy_id"] = str(clean.get("policy_id") or clean["sequence"])
    return Rule(**clean)


def parse_fortigate_config(path: Path, device="unknown", vdom="root") -> list[Rule]:
    rules=[]; in_policy=False; cur=None; seq=0
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line=raw.strip()
        if line == "config firewall policy": in_policy=True; continue
        if in_policy and line == "end":
            if cur: rules.append(cur)
            break
        if not in_policy: continue
        m=re.match(r"edit\s+\"?([^\"]+)\"?", line)
        if m:
            if cur: rules.append(cur)
            seq+=1; cur=Rule(device=device,vdom=vdom,policy_id=m.group(1),sequence=seq)
            continue
        if line == "next":
            if cur: rules.append(cur); cur=None
            continue
        m=re.match(r"set\s+(\S+)\s+(.+)", line)
        if m and cur:
            k,v=m.groups(); v=";".join(re.findall(r'"([^"]*)"|(\S+)', v)[0]) if False else v
            vals=re.findall(r'"([^"]*)"|(\S+)', v)
            val=";".join(a or b for a,b in vals)
            mapping={"policyid":"policy_id","srcintf":"srcintf","dstintf":"dstintf","srcaddr":"srcaddr",
                     "dstaddr":"dstaddr","service":"service","schedule":"schedule","action":"action",
                     "status":"status","name":"name","comments":"comments"}
            if k in mapping: setattr(cur,mapping[k],val)
    return rules


def token_is_any(t: str) -> bool: return t in {"all","any","0.0.0.0/0","::/0"}

def parse_net(t: str):
    try:
        if "-" in t:
            a,b=t.split("-",1); return (ipaddress.ip_address(a.strip()), ipaddress.ip_address(b.strip()))
        if re.match(r"^\d+\.\d+\.\d+\.\d+\s+\d+\.\d+\.\d+\.\d+$", t):
            a,m=t.split(); return ipaddress.ip_network(f"{a}/{m}", strict=False)
        return ipaddress.ip_network(t, strict=False)
    except ValueError: return None

def one_contains(a: str, b: str) -> bool:
    if token_is_any(a): return True
    if a == b: return True
    na,nb=parse_net(a),parse_net(b)
    if na is None or nb is None: return False
    if isinstance(na,tuple):
        if isinstance(nb,tuple): return na[0] <= nb[0] and na[1] >= nb[1]
        return na[0] <= nb.network_address and na[1] >= nb.broadcast_address
    if isinstance(nb,tuple): return nb[0] in na and nb[1] in na
    return nb.subnet_of(na)

def set_contains(A: tuple[str,...], B: tuple[str,...]) -> bool:
    return all(any(one_contains(a,b) for a in A) for b in B)

def overlaps(A: tuple[str,...], B: tuple[str,...]) -> bool:
    for a in A:
        for b in B:
            if one_contains(a,b) or one_contains(b,a): return True
    return False

def dimensions(r: Rule):
    n=r.normalized(); return [n[x] for x in ("srcintf","dstintf","srcaddr","dstaddr","service","schedule")]

def compare(a: Rule,b: Rule) -> dict|None:
    da,db=dimensions(a),dimensions(b)
    eq=all(x==y for x,y in zip(da,db)); a_contains=all(set_contains(x,y) for x,y in zip(da,db)); b_contains=all(set_contains(y,x) for x,y in zip(da,db)); ov=all(overlaps(x,y) for x,y in zip(da,db))
    same_action=a.action.lower()==b.action.lower()
    if eq and same_action: kind,severity="DUPLICATE","medium"
    elif a_contains and not same_action: kind,severity="SHADOWING","high"
    elif b_contains and not same_action: kind,severity="GENERALIZATION","review"
    elif a_contains and same_action: kind,severity="REDUNDANCY","low"
    elif ov and not same_action: kind,severity="CORRELATION","high"
    else: return None
    return {"device":a.device,"vdom":a.vdom,"rule_a":a.policy_id,"seq_a":a.sequence,"rule_b":b.policy_id,"seq_b":b.sequence,"type":kind,"severity":severity,"action_a":a.action,"action_b":b.action,"recommendation":recommendation(kind)}

def recommendation(kind):
    return {"DUPLICATE":"Validar y retirar una regla si no existe una justificación operativa.","SHADOWING":"Revisar orden y alcance; la regla inferior puede no ejecutarse para el tráfico cubierto.","GENERALIZATION":"Confirmar que la excepción específica y la regla general son intencionales.","REDUNDANCY":"Validar dependencias y retirar la regla contenida si es innecesaria.","CORRELATION":"Documentar o separar explícitamente el tráfico solapado y verificar el orden."}[kind]

def analyze(rules:list[Rule]):
    active=[r for r in rules if r.status.lower() not in {"disable","disabled"}]
    findings=[]
    groups={}
    for r in active: groups.setdefault((r.device,r.vdom),[]).append(r)
    for _,rs in groups.items():
        rs.sort(key=lambda x:x.sequence)
        for i,a in enumerate(rs):
            for b in rs[i+1:]:
                x=compare(a,b)
                if x: findings.append(x)
    return findings


def extract_inventory(path: Path) -> list[dict]:
    if path.suffix.lower()==".json": data=json.loads(path.read_text(encoding="utf-8-sig"))
    else:
        try:
            with tarfile.open(path,"r:*") as t:
                member=next((m for m in t.getmembers() if m.name.endswith("data.json") or m.name.endswith(".json")),None)
                if not member: raise ValueError("No se encontró data.json dentro del archivo")
                data=json.load(t.extractfile(member))
        except tarfile.ReadError as e: raise ValueError("El archivo no es un tar/gzip legible") from e
    records=[]
    def walk(x, adom=""):
        if isinstance(x,list):
            for y in x: walk(y,adom)
        elif isinstance(x,dict):
            new_adom=str(x.get("adom") or x.get("adom_name") or adom or "")
            keys={k.lower():v for k,v in x.items()}
            if any(k in keys for k in ("sn","serial","serial_number")) or ("hostname" in keys and any(k in keys for k in ("ip","ip_address","mgmt_ip"))):
                records.append({"adom":new_adom,"hostname":keys.get("hostname",keys.get("name","")),"serial":keys.get("sn",keys.get("serial",keys.get("serial_number",""))),"management_ip":keys.get("ip",keys.get("ip_address",keys.get("mgmt_ip",""))),"version":keys.get("version",keys.get("os_ver","")),"build":keys.get("build","")})
            for v in x.values(): walk(v,new_adom)
    walk(data)
    seen=set(); out=[]
    for r in records:
        key=tuple(str(r[k]) for k in sorted(r))
        if key not in seen: seen.add(key); out.append(r)
    return out


def write_csv(path:Path, rows:list[dict]):
    fields=list(rows[0].keys()) if rows else ["device","vdom","rule_a","rule_b","type","severity","recommendation"]
    with path.open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)

def write_xlsx(path:Path,rules,findings,inventory):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); wb.remove(wb.active)
    for title,rows in (("Resumen",[{"metrica":"Reglas analizadas","valor":len(rules)},{"metrica":"Hallazgos","valor":len(findings)},{"metrica":"Dispositivos inventario","valor":len(inventory)}]),("Hallazgos",findings),("Reglas",[asdict(r) for r in rules]),("Inventario",inventory)):
        ws=wb.create_sheet(title); fields=list(rows[0].keys()) if rows else ["sin_datos"]
        ws.append(fields)
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
        for row in rows: ws.append([row.get(f,"") for f in fields])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(c.value or "")) for c in col)+2))
    wb.save(path)

def main():
    p=argparse.ArgumentParser(description="Auditor estático de políticas Fortinet")
    p.add_argument("--policies",help="CSV, JSON o configuración FortiGate")
    p.add_argument("--format",choices=["auto","csv","json","fortigate"],default="auto")
    p.add_argument("--inventory",help="Exportación .dat/.tgz o data.json de FortiManager")
    p.add_argument("--output",default="audit_output",help="Directorio de salida")
    p.add_argument("--device",default="unknown"); p.add_argument("--vdom",default="root")
    a=p.parse_args(); out=Path(a.output); out.mkdir(parents=True,exist_ok=True)
    rules=[]
    if a.policies:
        path=Path(a.policies); fmt=a.format
        if fmt=="auto": fmt="csv" if path.suffix.lower()==".csv" else "json" if path.suffix.lower()==".json" else "fortigate"
        rules=read_csv(path) if fmt=="csv" else read_json(path) if fmt=="json" else parse_fortigate_config(path,a.device,a.vdom)
    inventory=extract_inventory(Path(a.inventory)) if a.inventory else []
    findings=analyze(rules)
    report={"metadata":{"rules":len(rules),"findings":len(findings),"inventory_devices":len(inventory),"sha256_policies":hashlib.sha256(Path(a.policies).read_bytes()).hexdigest() if a.policies else None},"findings":findings,"rules":[asdict(r) for r in rules],"inventory":inventory}
    (out/"audit_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    write_csv(out/"findings.csv",findings); write_csv(out/"inventory.csv",inventory)
    write_xlsx(out/"audit_report.xlsx",rules,findings,inventory)
    print(json.dumps(report["metadata"],ensure_ascii=False)); print(f"Resultados: {out.resolve()}")
if __name__=="__main__": main()
